#!/usr/bin/env python3
"""
py_pipeline_precheck.py — sanity check before running the pipeline.

Checks everything that would make py_stage00_normalize_jd_input.py /
run_pipeline.py fail partway through, before you actually run them:

  - CVTailor's own files are present where expected
  - Python is new enough (src/core/schemas.py and friends use `X | None`
    syntax, which needs 3.10+)
  - Required packages are installed
  - The active provider (LM Studio by default, or Claude/Gemini if
    explicitly requested) is actually reachable
  - Your input files exist AND actually open (not just "the file is
    there" — a resume.pdf that's actually 0 bytes, or a docx that's
    secretly corrupted, passes a naive exists() check and fails ingest
    three stages later)
  - The JD JSON validates against JDInput, auto-detecting harvester vs.
    canonical format the same way stage 0 does — but nothing gets
    written anywhere, this is read-only
  - applications/ and logs/ are writable

No API calls by default (free, instant). Pass --check-api-live if you
want one real, minimal request made to confirm the key actually works,
not just that it's set.

Usage:
    python src/checkpoints/py_pipeline_precheck.py \
        --jd-json /path/to/jd.json \
        --resume /path/to/resume.pdf \
        --cover-letter-sample /path/to/sample.txt \
        --resume-template /path/to/template.docx \
        [--cvtailor-root .] \
        [--check-api-live]
"""
import argparse
import importlib
import json
import os
import sys
from pathlib import Path

PASS, FAIL, WARN = "PASS", "FAIL", "WARN"
results: list[tuple[str, str, str]] = []


def check(label: str, ok: bool, detail: str = "", warn: bool = False) -> bool:
    status = PASS if ok else (WARN if warn else FAIL)
    results.append((status, label, detail))
    return ok


def check_python_version():
    ok = sys.version_info >= (3, 10)
    check(
        "Python version",
        ok,
        f"{sys.version.split()[0]} "
        + ("" if ok else "— need 3.10+ (src/core/schemas.py and friends use `X | None` syntax)"),
    )


def check_project_files(root: Path):
    core_dir = root / "src" / "core"
    for fname in ("schemas.py", "batch_common.py", "llm_dispatch.py", "io_utils.py"):
        check(f"project file: src/core/{fname}", (core_dir / fname).is_file(), str(core_dir / fname))
    check("project file: src/adapters/harvester_adapter.py",
          (root / "src" / "adapters" / "harvester_adapter.py").is_file(), "")
    check("project file: run_pipeline.py", (root / "run_pipeline.py").is_file(), "")
    check("project file: requirements.txt", (root / "requirements.txt").is_file(), "")

    stages_dir = root / "src" / "stages"
    if not check("src/stages/ directory", stages_dir.is_dir(), str(stages_dir)):
        return
    # Explicit expected filenames rather than glob()-discover-then-
    # startswith-match -- the latter failed on at least one real
    # Windows/Python 3.14 setup for reasons not fully root-caused (this
    # sandbox's Linux/Python 3.12 glob() found the same files fine, so
    # it's environment-specific, not a logic bug in the matching
    # itself). Checking each exact expected path removes the
    # indirection entirely rather than chasing a platform-specific
    # glob() quirk blind.
    expected_stage_files = {
        0: "py_stage00_normalize_jd_input.py",
        1: "py_stage02_ingest_cv_claims.py",
        2: "py_stage03_research_jd_company.py",
        3: "py_stage04_analyze_requirement_gaps.py",
        4: "py_stage05_propose_updates.py",
        5: "py_stage06_fact_check_updates.py",
        6: "py_stage08_agentic_update_review.py",
        7: "py_stage09_assemble_materials.py",
    }
    for i, fname in expected_stage_files.items():
        check(f"stage script: stage {i}", (stages_dir / fname).is_file(), fname)


def _resolve_provider() -> str:
    sys.path.insert(0, str(Path(".").resolve() / "src" / "core"))
    try:
        from llm_dispatch import default_provider
        return default_provider()
    except Exception:
        return os.environ.get("CVTAILOR_LLM_PROVIDER", "lmstudio")


def check_dependencies():
    for module, pip_name in [
        ("openai", "openai"), ("pydantic", "pydantic"),
        ("docx", "python-docx"), ("pdfplumber", "pdfplumber"),
    ]:
        try:
            importlib.import_module(module)
            note = "" if module != "openai" else "(required for the default lmstudio provider)"
            check(f"dependency: {pip_name}", True, note)
        except ImportError:
            check(f"dependency: {pip_name}", False,
                  f"pip install --break-system-packages {pip_name}")

    provider = _resolve_provider()

    for module, pip_name, flag in [("anthropic", "anthropic", "claude"), ("google.genai", "google-genai", "gemini")]:
        try:
            importlib.import_module(module)
            check(f"dependency: {pip_name}", True,
                  "" if provider == flag else f"(not needed unless --provider {flag})")
        except ImportError:
            if provider == flag:
                check(f"dependency: {pip_name}", False,
                      f"pip install --break-system-packages {pip_name} — required for --provider {flag}")
            else:
                check(f"dependency: {pip_name}", False,
                      f"not installed, but not needed unless --provider {flag}", warn=True)


def check_api_key():
    provider = _resolve_provider()

    if provider == "gemini":
        key = os.environ.get("GEMINI_API_KEY", "")
        check("GEMINI_API_KEY set (provider=gemini)", bool(key),
              "" if key else "set it before running any stage that calls Gemini")
    elif provider == "claude":
        key = os.environ.get("ANTHROPIC_API_KEY", "")
        check("ANTHROPIC_API_KEY set (provider=claude)", bool(key),
              "" if key else "set it before running any stage that calls Claude")
    else:  # lmstudio -- no API key needed, but the local server needs to actually be reachable
        sys.path.insert(0, str(Path(".").resolve() / "src" / "core"))
        try:
            from llm_dispatch import LMSTUDIO_BASE_URL
        except Exception:
            LMSTUDIO_BASE_URL = os.environ.get("LMSTUDIO_BASE_URL", "http://localhost:1234/v1")
        try:
            import openai
            client = openai.OpenAI(base_url=LMSTUDIO_BASE_URL, api_key="lm-studio")
            models = client.models.list()
            if models.data:
                names = ", ".join(m.id for m in models.data[:3])
                check(f"LM Studio reachable at {LMSTUDIO_BASE_URL} (provider=lmstudio)", True,
                      f"model(s) loaded: {names}")
            else:
                check(f"LM Studio reachable at {LMSTUDIO_BASE_URL} (provider=lmstudio)", False,
                      "server responded but no model is loaded — load one in LM Studio's UI")
        except Exception as e:
            check(f"LM Studio reachable at {LMSTUDIO_BASE_URL} (provider=lmstudio)", False,
                  f"{type(e).__name__}: {e} — start the server in LM Studio (Developer tab -> "
                  f"Start Server), or use --provider claude/gemini instead")


def check_jd_json(root: Path, jd_json_path: str, label: str | None = None):
    label = label or "JD JSON"
    p = Path(jd_json_path)
    if not check(f"{label} file exists", p.is_file(), jd_json_path):
        return
    try:
        raw = json.loads(p.read_text(encoding="utf-8"))
    except Exception as e:
        check(f"{label} is valid JSON", False, str(e))
        return
    check(f"{label} is valid JSON", True)

    sys.path.insert(0, str(root / "src" / "core"))
    sys.path.insert(0, str(root / "src" / "adapters"))
    try:
        from schemas import JDInput
        from harvester_adapter import is_harvester_format, normalize_harvester_json
    except Exception as e:
        check("can import src/core/schemas.py / src/adapters/harvester_adapter.py", False, str(e))
        return
    check("can import src/core/schemas.py / src/adapters/harvester_adapter.py", True)

    fmt = "job_data_harvester" if is_harvester_format(raw) else "canonical_jdinput"
    try:
        normalized = normalize_harvester_json(raw) if fmt == "job_data_harvester" else raw
        jd = JDInput.model_validate(normalized)
        check(f"{label} validates against JDInput ({fmt})", True,
              f"company={jd.company!r} role={jd.role_title!r}")
    except Exception as e:
        check(f"{label} validates against JDInput ({fmt})", False, str(e))


def check_resume(resume_path: str):
    p = Path(resume_path)
    if not check("resume file exists", p.is_file(), resume_path):
        return
    if not check("resume file non-empty", p.stat().st_size > 0, f"{p.stat().st_size} bytes"):
        return

    suffix = p.suffix.lower()
    if suffix == ".pdf":
        try:
            import pdfplumber
            with pdfplumber.open(p) as pdf:
                n = len(pdf.pages)
            check("resume PDF opens and has pages", n > 0, f"{n} page(s)")
        except Exception as e:
            check("resume PDF opens and has pages", False, str(e))
    elif suffix == ".docx":
        try:
            import docx
            d = docx.Document(p)
            check("resume docx opens", True, f"{len(d.paragraphs)} paragraph(s)")
        except Exception as e:
            check("resume docx opens", False, str(e))
    else:
        check("resume file type", False,
              f"unexpected extension {suffix!r} — py_stage02_ingest_cv_claims.py expects .pdf or .docx", warn=True)


def check_cover_letter_sample(path: str):
    p = Path(path)
    if not check("cover letter sample exists", p.is_file(), path):
        return
    try:
        text = p.read_text(encoding="utf-8")
        check("cover letter sample readable as text", len(text.strip()) > 0, f"{len(text)} chars")
    except Exception as e:
        check("cover letter sample readable as text", False, str(e))


def check_resume_template(path: str):
    p = Path(path)
    if not check("resume template exists", p.is_file(), path):
        return
    try:
        import docx
        d = docx.Document(p)
        check("resume template docx opens", True,
              f"{len(d.paragraphs)} paragraph(s), {len(d.tables)} table(s)")
    except Exception as e:
        check("resume template docx opens", False, str(e))


def check_writable_dirs(root: Path):
    for dirname in ("applications", "logs"):
        d = root / dirname
        try:
            d.mkdir(parents=True, exist_ok=True)
            probe = d / ".cvtailor_write_test"
            probe.write_text("ok")
            probe.unlink()
            check(f"{dirname}/ writable", True, str(d))
        except Exception as e:
            check(f"{dirname}/ writable", False, str(e))


def check_batch_config(root: Path):
    sys.path.insert(0, str(root / "src" / "core"))
    from batch_common import (
        COVER_LETTER_SAMPLE_DIR, APPLICATIONS_ROOT,
        discover_resume_variants, discover_resume_template_variants,
        discover_single_file, load_applicant_info, app_dirs, find_jd_input,
    )

    try:
        variants = discover_resume_variants()
        label = ", ".join(f"{name} ({path.name})" for name, path in sorted(variants.items()))
        check("config/resume/ has at least one resume", True,
              f"{len(variants)} variant(s): {label}")
    except Exception as e:
        check("config/resume/ has at least one resume", False, str(e))
        variants = {}

    try:
        templates = discover_resume_template_variants()
        label = ", ".join(f"{name} ({path.name})" for name, path in sorted(templates.items()))
        check("config/resume_template/ has at least one template", True,
              f"{len(templates)} variant(s): {label}")
    except Exception as e:
        check("config/resume_template/ has at least one template", False, str(e))

    try:
        sample = discover_single_file(COVER_LETTER_SAMPLE_DIR, (".txt",), "cover letter sample")
        check("config/cover_letter_sample/ has exactly one file", True, sample.name)
    except Exception as e:
        check("config/cover_letter_sample/ has exactly one file", False, str(e))

    try:
        applicant = load_applicant_info()
        check("config/applicant_info.json valid", True, f"name={applicant['name']!r}")
    except Exception as e:
        check("config/applicant_info.json valid", False, str(e))

    # applications/ is both the queue and the drop zone now -- no
    # separate config/job_descriptions/intake/ to check. Count raw
    # (not-yet-validated) files at the root plus already-queued
    # jd_input.json folders, so this still tells you whether there's
    # actually anything for run_pipeline.py to do.
    APPLICATIONS_ROOT.mkdir(parents=True, exist_ok=True)
    raw_at_root = sorted(p for p in APPLICATIONS_ROOT.iterdir() if p.is_file() and p.suffix.lower() == ".json")
    queued_dirs = [d for d in app_dirs() if find_jd_input(d) is not None]
    total = len(raw_at_root) + len(queued_dirs)
    if not check("applications/ has at least one JD (raw or already queued)",
                  total > 0, f"{len(raw_at_root)} raw file(s) at root, "
                             f"{len(queued_dirs)} already-queued folder(s)"):
        return
    for jd_file in raw_at_root:
        check_jd_json(root, str(jd_file), label=f"applications/{jd_file.name}")


def check_docx2pdf():
    try:
        importlib.import_module("docx2pdf")
        check("dependency: docx2pdf (optional)", True,
              "final output will be .pdf; falls back to .docx if this ever fails at runtime")
    except ImportError:
        check("dependency: docx2pdf (optional)", False,
              "pip install --break-system-packages docx2pdf (Windows/Mac + MS Word only) — "
              "output will be .docx instead of .pdf without it", warn=True)


def check_api_live():
    provider = _resolve_provider()

    if provider == "gemini":
        try:
            from google import genai
            client = genai.Client(api_key=os.environ.get("GEMINI_API_KEY"))
            client.models.generate_content(model="gemini-3.5-flash-lite", contents="hi")
            check("live API call succeeds (gemini)", True, "reached Gemini API with current GEMINI_API_KEY")
        except Exception as e:
            check("live API call succeeds (gemini)", False, str(e))
        return

    if provider == "lmstudio":
        try:
            import openai
            sys.path.insert(0, str(Path(".").resolve() / "src" / "core"))
            try:
                from llm_dispatch import LMSTUDIO_BASE_URL, _lmstudio_resolve_model
                client = openai.OpenAI(base_url=LMSTUDIO_BASE_URL, api_key="lm-studio")
                model = _lmstudio_resolve_model(client)
            except ImportError:
                LMSTUDIO_BASE_URL = os.environ.get("LMSTUDIO_BASE_URL", "http://localhost:1234/v1")
                client = openai.OpenAI(base_url=LMSTUDIO_BASE_URL, api_key="lm-studio")
                model = client.models.list().data[0].id
            client.chat.completions.create(
                model=model, max_tokens=1, messages=[{"role": "user", "content": "hi"}],
            )
            check("live API call succeeds (lmstudio)", True, f"reached LM Studio, model={model!r}")
        except Exception as e:
            check("live API call succeeds (lmstudio)", False, str(e))
        return

    try:
        import anthropic
        client = anthropic.Anthropic()
        client.messages.create(
            model="claude-haiku-4-5-20251001",
            max_tokens=1,
            messages=[{"role": "user", "content": "hi"}],
        )
        check("live API call succeeds (claude)", True, "reached Anthropic API with current ANTHROPIC_API_KEY")
    except Exception as e:
        check("live API call succeeds (claude)", False, str(e))


def print_report() -> int:
    width = max(len(label) for _, label, _ in results) + 2
    n_fail = sum(1 for s, _, _ in results if s == FAIL)
    n_warn = sum(1 for s, _, _ in results if s == WARN)
    print()
    for status, label, detail in results:
        marker = {"PASS": "OK ", "FAIL": "!! ", "WARN": " ? "}[status]
        print(f"  [{marker}] {label:<{width}} {detail}")
    print()
    print(f"{len(results)} checks — {n_fail} failed, {n_warn} warning(s)")
    print("Ready — intake should run cleanly." if n_fail == 0
          else "Fix the failed items above before running run_pipeline.py.")
    return n_fail


def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                  formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--jd-json", default=None,
                     help="Explicit single-JD mode. Omit all four of --jd-json/--resume/"
                          "--cover-letter-sample/--resume-template to check the config/ "
                          "batch structure instead.")
    ap.add_argument("--resume", default=None)
    ap.add_argument("--cover-letter-sample", default=None)
    ap.add_argument("--resume-template", default=None)
    ap.add_argument("--cvtailor-root", default=".")
    ap.add_argument("--check-api-live", action="store_true",
                     help="Make one real, minimal API call to confirm the key actually works. "
                          "Off by default — costs a fraction of a cent, but that's still not zero.")
    args = ap.parse_args()

    root = Path(args.cvtailor_root).resolve()
    explicit_args = [args.jd_json, args.resume, args.cover_letter_sample, args.resume_template]
    batch_mode = not any(explicit_args)

    if not batch_mode and not all(explicit_args):
        print("Provide all four of --jd-json/--resume/--cover-letter-sample/--resume-template "
              "for single-file mode, or none of them to check the config/ batch structure.",
              file=sys.stderr)
        sys.exit(2)

    check_python_version()
    check_project_files(root)
    check_dependencies()
    check_docx2pdf()
    check_api_key()

    if batch_mode:
        print("(no --jd-json/--resume/etc. given — checking config/ batch structure)", file=sys.stderr)
        check_batch_config(root)
    else:
        check_jd_json(root, args.jd_json)
        check_resume(args.resume)
        check_cover_letter_sample(args.cover_letter_sample)
        check_resume_template(args.resume_template)

    check_writable_dirs(root)
    if args.check_api_live:
        check_api_live()

    n_fail = print_report()
    sys.exit(1 if n_fail else 0)


if __name__ == "__main__":
    main()
